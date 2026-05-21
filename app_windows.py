"""
app_windows.py — LoginWin and AppWin (main application with all pages).
"""

import os
import re
import csv
import sqlite3
import datetime
import shutil
import webbrowser
import tkinter as tk
from tkinter import ttk, messagebox, filedialog, simpledialog

from config import (W, BLK, G1, G2, G3, G4, G5, ACC, ACC2,
                    RLT, GLT, YLT, SEL, F, IS_DARK)
from database import con, hp, now, tod, init_db, next_no, audit, get_setting, can_edit
from ui_helpers import setup_styles, mk_btn, mk_tree, divider, stat_box, open_file
from dialogs import HistWin, PrintWin, UserDlg, UserEditDlg, ChangePwDlg


# ══════════════════════════════════════════════════════════
#  LOGIN WINDOW
# ══════════════════════════════════════════════════════════
class LoginWin(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Document Management System")
        self.configure(bg=W)
        self.state("zoomed")
        init_db()
        setup_styles()
        self._build()

    def _build(self):
        container = tk.Frame(self, bg=W)
        container.place(relx=0.5, rely=0.45, anchor=tk.CENTER)

        banner = tk.Frame(container, bg=ACC, width=450, height=120)
        banner.pack(fill=tk.X)
        banner.pack_propagate(False)
        tk.Label(banner, text="LOGIN", font=(F, 32, "bold"), bg=ACC, fg="#ffffff").pack(pady=(35, 0))

        form = tk.Frame(container, bg=W, bd=1, relief=tk.SOLID)
        form.pack(fill=tk.X, expand=True, ipadx=40, ipady=30)

        tk.Label(form, text="Username", font=(F, 10, "bold"), bg=W, fg=BLK).pack(anchor="w")
        u_wrap = tk.Frame(form, bg=G3, bd=0)
        u_wrap.pack(fill=tk.X, pady=(4, 16))
        self._u = tk.Entry(u_wrap, width=38, bg=G2, fg=BLK, insertbackground=BLK,
                           relief=tk.FLAT, font=(F, 11), bd=10)
        self._u.pack(fill=tk.X)

        tk.Label(form, text="Password", font=(F, 10, "bold"), bg=W, fg=BLK).pack(anchor="w")
        pw_fr = tk.Frame(form, bg=G3)
        pw_fr.pack(fill=tk.X, pady=(4, 6))
        self._pw = tk.Entry(pw_fr, show="*", width=36, bg=G2, fg=BLK, insertbackground=BLK,
                            relief=tk.FLAT, font=(F, 11), bd=10)
        self._pw.pack(side=tk.LEFT, fill=tk.X, expand=True)

        self._show = False
        self._eye = tk.Button(pw_fr, text="Show", command=self._toggle_pw,
                              bg=G2, fg=G4, relief=tk.FLAT, font=(F, 8),
                              cursor="hand2", bd=0, padx=10, pady=0)
        self._eye.pack(side=tk.RIGHT)

        self._msg = tk.StringVar()
        tk.Label(form, textvariable=self._msg, font=(F, 9), bg=W, fg="#b91c1c").pack(anchor="w", pady=(4, 0))
        mk_btn(form, "LOGIN", self._login, style="p", py=12, px=0).pack(fill=tk.X, pady=(20, 0))

        self.bind("<Return>", lambda e: self._login())
        self._u.focus_set()

    def _toggle_pw(self):
        self._show = not self._show
        self._pw.config(show="" if self._show else "*")
        self._eye.config(text="Hide" if self._show else "Show")

    def _login(self):
        u = self._u.get().strip()
        p = self._pw.get()
        if not u or not p:
            self._msg.set("Username and Password are required!"); return
        c = con(); cur = c.cursor()
        cur.execute("SELECT * FROM users WHERE username=? AND password=? AND is_active=1", (u, hp(p)))
        row = cur.fetchone()
        if row:
            c.execute("UPDATE users SET last_login=? WHERE id=?", (now(), row[0]))
            c.commit()
        c.close()

        if row:
            self._msg.set("")
            user = {
                "id":        row[0],
                "username":  row[1],
                "full_name": row[3] or row[1],
                "role":      row[7],
            }
            self.withdraw()
            AppWin(self, user)
        else:
            self._msg.set("Invalid username or password!")
            self._pw.delete(0, tk.END)
            self._pw.focus_set()


# ══════════════════════════════════════════════════════════
#  MAIN APPLICATION WINDOW
# ══════════════════════════════════════════════════════════
class AppWin(tk.Toplevel):
    def __init__(self, login, user):
        super().__init__()
        self.login = login
        self.user  = user
        self._sel  = None
        self.title("Document Management System")
        self.configure(bg=G1)
        self.protocol("WM_DELETE_WINDOW", self._quit)
        self.minsize(1024, 600)
        setup_styles()
        self._build()
        self._go("dashboard")
        self.state("zoomed")

    def _build(self):
        self._topbar()
        body = tk.Frame(self, bg=G1)
        body.pack(fill=tk.BOTH, expand=True)
        self._sidebar(body)
        self.pane = tk.Frame(body, bg=G1)
        self.pane.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

    # ────────────────────────────────────────────
    #  TOP BAR
    # ────────────────────────────────────────────
    def _topbar(self):
        bar = tk.Frame(self, bg=ACC, height=56)
        bar.pack(fill=tk.X); bar.pack_propagate(False)
        tk.Label(bar, text="  DOCUMENT MANAGEMENT SYSTEM",
                 font=(F, 13, "bold"), bg=ACC, fg="#ffffff").pack(side=tk.LEFT)

        rht = tk.Frame(bar, bg=ACC)
        rht.pack(side=tk.RIGHT, padx=12)
        self._ab = None

        nf = tk.Frame(rht, bg=ACC, cursor="hand2")
        nf.pack(side=tk.LEFT, padx=(0, 8))
        name_txt = self.user["full_name"] or self.user["username"]
        self._name_lbl = tk.Label(nf, text=name_txt + " ▼",
                                  font=(F, 10, "bold"), bg=ACC, fg="#ffffff", cursor="hand2")
        self._name_lbl.pack(anchor="e")
        role_colors = {"Admin": "#f59e0b", "Clerk": "#38bdf8", "User": "#9dd8be", "Manager": "#a78bfa"}
        role_fg = role_colors.get(self.user["role"], "#9dd8be")
        tk.Label(nf, text=self.user["role"], font=(F, 8, "bold"),
                 bg=ACC, fg=role_fg, cursor="hand2").pack(anchor="e")

        av = tk.Canvas(rht, width=36, height=36, bg=ACC, highlightthickness=0, cursor="hand2")
        av.pack(side=tk.LEFT, pady=10)
        av.create_oval(1, 1, 35, 35, fill="#ffffff", outline="#ffffff", width=1)
        ini = "".join(w[0].upper() for w in (self.user["full_name"] or self.user["username"]).split()[:2])
        av.create_text(18, 18, text=ini, font=(F, 12, "bold"), fill=ACC)

        av.bind("<Button-1>", self._show_user_menu)
        self._name_lbl.bind("<Button-1>", self._show_user_menu)
        nf.bind("<Button-1>", self._show_user_menu)
        self._upd_alerts()

    def _show_user_menu(self, event):
        menu = tk.Menu(self, tearoff=0, font=(F, 10), bg=W, fg=BLK)
        if self.user["role"] == "Admin":
            menu.add_command(label="➕ Add New User",
                             command=lambda: UserDlg(self, None, lambda: None))
            menu.add_separator()
        menu.add_command(label="🔑 Change Password",
                         command=lambda: ChangePwDlg(self, self.user["id"]))
        menu.post(event.x_root, event.y_root + 15)

    # ────────────────────────────────────────────
    #  SIDEBAR
    # ────────────────────────────────────────────
    def _sidebar(self, parent):
        sb = tk.Frame(parent, bg=W, width=210, bd=0)
        sb.pack(side=tk.LEFT, fill=tk.Y); sb.pack_propagate(False)
        tk.Frame(sb, bg=G3, width=1).pack(side=tk.RIGHT, fill=tk.Y)

        self._nbtn = {}
        nav = [
            ("dashboard", "  Dashboard",         "Admin,Manager,Clerk,User"),
            ("docs",      "  All Documents",     "Admin,Manager,Clerk,User"),
            ("new_in",    "⬇  New Inward",       "Admin,Manager,Clerk"),
            ("new_out",   "⬆  New Outward",      "Admin,Manager,Clerk"),
            ("user_mgmt", "👥  User Management", "Admin"),
        ]
        tk.Label(sb, text="NAVIGATION", font=(F, 8, "bold"), bg=W, fg=G4).pack(
            anchor="w", padx=16, pady=(14, 6))

        for key, label, roles in nav:
            if self.user["role"] not in roles.split(","):
                continue
            b = tk.Button(sb, text=label, command=lambda k=key: self._go(k),
                          bg=W, fg=BLK, relief=tk.FLAT, anchor="w",
                          font=(F, 10), padx=16, pady=9, cursor="hand2",
                          activebackground=GLT, activeforeground=BLK, bd=0)
            b.pack(fill=tk.X)
            self._nbtn[key] = b

        tk.Frame(sb, bg=W).pack(fill=tk.Y, expand=True)
        tk.Frame(sb, bg=G3, height=1).pack(fill=tk.X)
        mk_btn(sb, "Logout", self._logout, style="d", py=10, px=16).pack(
            fill=tk.X, padx=10, pady=10)

    def _go(self, key):
        for w in self.pane.winfo_children(): w.destroy()
        self._sel = None
        for k, b in self._nbtn.items():
            b.config(bg=ACC if k == key else W,
                     fg="#ffffff" if k == key else BLK,
                     font=(F, 10, "bold" if k == key else "normal"))
        {
            "dashboard": self._pg_dashboard,
            "docs":      self._pg_docs,
            "new_in":    lambda: self._pg_form("INWARD"),
            "new_out":   lambda: self._pg_form("OUTWARD"),
            "user_mgmt": self._pg_user_mgmt,
        }.get(key, self._pg_dashboard)()

    def _page(self, title, sub=""):
        wrap = tk.Frame(self.pane, bg=G1)
        wrap.pack(fill=tk.BOTH, expand=True)
        ph = tk.Frame(wrap, bg=W); ph.pack(fill=tk.X)
        tk.Frame(ph, bg=G3, height=1).pack(fill=tk.X, side=tk.BOTTOM)
        ih = tk.Frame(ph, bg=W); ih.pack(fill=tk.X, padx=24, pady=14)
        tk.Label(ih, text=title, font=(F, 16, "bold"), bg=W, fg=BLK).pack(side=tk.LEFT)
        if sub:
            tk.Label(ih, text="  " + sub, font=(F, 9), bg=W, fg=G4).pack(side=tk.LEFT, pady=(4, 0))
        body = tk.Frame(wrap, bg=W)
        body.pack(fill=tk.BOTH, expand=True, padx=24, pady=10)
        return body, ih

    # ════════════════════════════════════════════
    #  DASHBOARD
    # ════════════════════════════════════════════
    def _pg_dashboard(self):
        body, _ = self._page("Dashboard",
                             f"Welcome back, {self.user['full_name'] or self.user['username']}")
        c = con(); cur = c.cursor()
        def q(sql): cur.execute(sql); return cur.fetchone()[0]
        total   = q("SELECT COUNT(*) FROM documents")
        inward  = q("SELECT COUNT(*) FROM documents WHERE doc_type='INWARD'")
        outward = q("SELECT COUNT(*) FROM documents WHERE doc_type='OUTWARD'")
        pending = q("SELECT COUNT(*) FROM documents WHERE status='Pending'")
        c.close()

        sg = tk.Frame(body, bg=W); sg.pack(fill=tk.X, pady=(10, 30))
        for col, (title, val, tint) in enumerate([
            ("Total Documents", total, G1), ("Inward", inward, GLT),
            ("Outward", outward, RLT),      ("Pending", pending, YLT),
        ]):
            stat_box(sg, col, title, val, tint)

        divider(body, 15)
        tk.Label(body, text="Recent Documents", font=(F, 12, "bold"),
                 bg=W, fg=BLK).pack(anchor="w", pady=(10, 6))

        row_limit = int(get_setting("dash_rows", "10"))
        frm, tv = mk_tree(body, ("Doc No", "Type", "Subject", "Sender / Receiver", "Date", "Status"),
                          (120, 80, 280, 200, 100, 100), height=row_limit)
        frm.pack(fill=tk.BOTH, expand=True)

        c = con(); cur = c.cursor()
        cur.execute(f"SELECT doc_no, doc_type, subject, sender_receiver, date_entry, status "
                    f"FROM documents ORDER BY doc_no ASC LIMIT {row_limit}")
        for row in cur.fetchall():
            tv.insert("", "end", values=row, tags=("IN" if row[1] == "INWARD" else "OUT",))
        c.close()

    # ════════════════════════════════════════════
    #  ALL DOCUMENTS
    # ════════════════════════════════════════════
    def _pg_docs(self):
        body, hdr = self._page("All Documents")
        if can_edit(self.user["role"]):
            mk_btn(hdr, "+ New Outward", lambda: self._go("new_out"), style="d", py=5).pack(side=tk.RIGHT, padx=(4, 0))
            mk_btn(hdr, "+ New Inward",  lambda: self._go("new_in"),  style="p", py=5).pack(side=tk.RIGHT, padx=(4, 0))

        fb = tk.Frame(body, bg=G1, bd=1, relief=tk.SOLID); fb.pack(fill=tk.X, pady=(0, 10))
        fbr = tk.Frame(fb, bg=G1); fbr.pack(fill=tk.X, padx=10, pady=7)

        tk.Label(fbr, text="Search:", font=(F, 9, "bold"), bg=G1, fg=BLK).pack(side=tk.LEFT)
        self._ds = tk.StringVar()
        self._ds.trace("w", lambda *a: self._ld())
        sw = tk.Frame(fbr, bg=G3); sw.pack(side=tk.LEFT, padx=(4, 10))
        tk.Entry(sw, textvariable=self._ds, width=20, bg=G2, fg=BLK,
                 insertbackground=BLK, relief=tk.FLAT, font=(F, 9), bd=6).pack()

        self._ft  = tk.StringVar(value="ALL")
        self._fst = tk.StringVar(value="ALL")
        self._fp  = tk.StringVar(value="ALL")
        self._fd  = tk.StringVar(value="ALL")

        for lbl, var, vals in [
            ("Type:",     self._ft,  ["ALL", "INWARD", "OUTWARD"]),
            ("Status:",   self._fst, ["ALL", "Pending", "In-Progress", "Completed", "Forwarded", "Archived"]),
            ("Priority:", self._fp,  ["ALL", "Urgent", "Normal", "Low"]),
            ("Dept:",     self._fd,  ["ALL", "Administration", "Finance", "HR", "IT", "Legal",
                                       "Operations", "Procurement", "Sales", "Other"]),
        ]:
            tk.Label(fbr, text=lbl, font=(F, 9, "bold"), bg=G1, fg=BLK).pack(side=tk.LEFT)
            cb = ttk.Combobox(fbr, textvariable=var, values=vals, width=11, font=(F, 9), state="readonly")
            cb.pack(side=tk.LEFT, padx=(2, 8), ipady=3)
            cb.bind("<<ComboboxSelected>>", lambda e: self._ld())

        cols   = ("Doc No", "Type", "Subject", "Sender/Receiver", "Dept",
                  "Entry Date", "Due Date", "Priority", "Status", "Ref", "File")
        widths = [120, 74, 210, 155, 82, 88, 88, 70, 86, 70, 36]

        ab = tk.Frame(body, bg=G1); ab.pack(side=tk.BOTTOM, fill=tk.X)
        self._sl = tk.Label(ab, text="", font=(F, 9), bg=ACC, fg="#ffffff", anchor="w")
        self._sl.pack(side=tk.LEFT, fill=tk.X, expand=True, ipady=5, padx=6)
        mk_btn(ab, "Open File", self._act_open, style="s", py=4, px=10, side=tk.RIGHT)
        if can_edit(self.user["role"]):
            mk_btn(ab, "Delete", self._act_delete, style="d", py=4, px=10, side=tk.RIGHT)
            mk_btn(ab, "Edit",   self._act_edit,   style="p", py=4, px=10, side=tk.RIGHT)

        tf, self._tv = mk_tree(body, cols, widths)
        tf.pack(fill=tk.BOTH, expand=True, pady=(0, 4))
        self._tv.bind("<<TreeviewSelect>>", self._ds_sel)
        self._tv.bind("<Double-1>",          self._ds_dbl)
        self._ld()

    def _ld(self):
        if not hasattr(self, "_tv"): return
        for r in self._tv.get_children(): self._tv.delete(r)
        srch = getattr(self, "_ds", tk.StringVar()).get().strip().lower()
        ft   = getattr(self, "_ft",  tk.StringVar(value="ALL")).get()
        fst  = getattr(self, "_fst", tk.StringVar(value="ALL")).get()
        fp   = getattr(self, "_fp",  tk.StringVar(value="ALL")).get()
        fd   = getattr(self, "_fd",  tk.StringVar(value="ALL")).get()
        c = con(); cur = c.cursor()
        q = ("SELECT doc_no,doc_type,subject,sender_receiver,department,"
             "date_entry,due_date,priority,status,ref_no,file_path,id FROM documents")
        params = []; conds = []
        if ft  != "ALL": conds.append("doc_type=?");   params.append(ft)
        if fst != "ALL": conds.append("status=?");     params.append(fst)
        if fp  != "ALL": conds.append("priority=?");   params.append(fp)
        if fd  != "ALL": conds.append("department=?"); params.append(fd)
        if srch:
            conds.append("(LOWER(subject) LIKE ? OR LOWER(sender_receiver) LIKE ?"
                         " OR LOWER(doc_no) LIKE ? OR LOWER(tags) LIKE ?)")
            params += [f"%{srch}%"] * 4
        if conds: q += " WHERE " + " AND ".join(conds)
        q += " ORDER BY doc_no ASC"
        cur.execute(q, params); rows = cur.fetchall(); c.close()

        for row in rows:
            base_tag = "IN" if row[1] == "INWARD" else "OUT"
            tags = (base_tag, "urg") if (row[7] == "Urgent" and row[8] not in ("Completed", "Archived")) else (base_tag,)
            icon = "✓" if row[10] and os.path.exists(row[10]) else ""
            self._tv.insert("", "end", iid=str(row[11]), values=row[:10] + (icon,), tags=tags)

        ic = sum(1 for r in rows if r[1] == "INWARD")
        oc = sum(1 for r in rows if r[1] == "OUTWARD")
        pc = sum(1 for r in rows if r[8] == "Pending")
        if hasattr(self, "_sl"):
            self._sl.config(text=f"  Total: {len(rows)}  |  Inward: {ic}  |  Outward: {oc}  |  Pending: {pc}")

    def _ds_sel(self, _):
        sel = self._tv.selection()
        if sel: self._sel = int(sel[0])

    def _ds_dbl(self, _):
        if self._sel: self._act_open()

    def _need_sel(self):
        if not self._sel:
            messagebox.showwarning("Select", "Please select a record first!")
            return False
        return True

    def _row(self):
        if not self._need_sel(): return None
        c = con(); cur = c.cursor()
        cur.execute("SELECT * FROM documents WHERE id=?", (self._sel,))
        row = cur.fetchone(); c.close(); return row

    def _act_edit(self):
        if not can_edit(self.user["role"]):
            messagebox.showwarning("Access Denied", "You do not have permission to edit documents."); return
        row = self._row()
        if row: self._pg_form(edit_row=row)

    def _act_delete(self):
        if not can_edit(self.user["role"]):
            messagebox.showwarning("Access Denied", "You do not have permission to delete documents."); return
        if not self._need_sel(): return
        c = con(); cur = c.cursor()
        cur.execute("SELECT doc_no FROM documents WHERE id=?", (self._sel,))
        dn = cur.fetchone(); c.close()
        if not messagebox.askyesno("Delete", f"Delete document '{dn[0] if dn else '?'}'?"): return
        c = con()
        c.execute("DELETE FROM documents WHERE id=?",      (self._sel,))
        c.execute("DELETE FROM doc_history WHERE doc_id=?", (self._sel,))
        c.commit(); c.close()
        self._sel = None; self._ld()

    def _act_history(self):
        if not self._need_sel(): return
        HistWin(self, self._sel)

    def _act_print(self):
        row = self._row()
        if not row: return
        c = con(); cur = c.cursor()
        cur.execute("SELECT * FROM documents WHERE id=?", (self._sel,))
        hdrs = [d[0] for d in cur.description]; c.close()
        PrintWin(self, row, hdrs, self.user)

    def _act_open(self):
        if not self._need_sel(): return
        c = con(); cur = c.cursor()
        cur.execute("SELECT file_path FROM documents WHERE id=?", (self._sel,))
        row = cur.fetchone(); c.close()
        if row and row[0]: open_file(row[0])
        else: messagebox.showinfo("No File", "No file attached to this document.")

    # ════════════════════════════════════════════
    #  DOCUMENT FORM (New / Edit)
    # ════════════════════════════════════════════
    def _pg_form(self, default_type="INWARD", edit_row=None):
        for w in self.pane.winfo_children(): w.destroy()
        title = "Edit Document" if edit_row else f"New {default_type.title()} Document"
        body, hdr = self._page(title)
        for k, b in self._nbtn.items(): b.config(bg=W, fg=BLK, font=(F, 10, "normal"))

        fc = tk.Frame(body, bg=W, bd=1, relief=tk.SOLID)
        fc.pack(fill=tk.X, pady=(0, 14))
        fc.columnconfigure(1, weight=1); fc.columnconfigure(3, weight=1)

        def lbl(text, r, c=0):
            tk.Label(fc, text=text, font=(F, 9, "bold"), bg=W, fg=BLK).grid(
                row=r, column=c, sticky="w", padx=(14, 4), pady=(10, 1))

        def ent(r, c=1, w=26, val="", show=""):
            wrap = tk.Frame(fc, bg=G3)
            wrap.grid(row=r, column=c, padx=(2, 12), pady=(1, 6), sticky="ew")
            e = tk.Entry(wrap, width=w, show=show, bg=G2, fg=BLK, insertbackground=BLK,
                         relief=tk.FLAT, font=(F, 10), bd=7)
            e.pack(fill=tk.X)
            if val: e.insert(0, val)
            return e

        def cbo(r, c=1, vals=[], val="", w=24):
            v = tk.StringVar(value=val)
            wrap = tk.Frame(fc, bg=G3)
            wrap.grid(row=r, column=c, padx=(2, 12), pady=(1, 6), sticky="ew")
            cb = ttk.Combobox(wrap, textvariable=v, values=vals, width=w,
                              font=(F, 10), state="readonly")
            cb.pack(fill=tk.X, ipady=5)
            return v

        lbl("Document Type *", 0)
        v_type = tk.StringVar(value=edit_row[1] if edit_row else default_type)
        tf = tk.Frame(fc, bg=W)
        tf.grid(row=0, column=1, columnspan=3, sticky="w", padx=(2, 12), pady=(10, 6))
        for txt, val in [("  INWARD  (Incoming)", "INWARD"), ("  OUTWARD  (Outgoing)", "OUTWARD")]:
            tk.Radiobutton(tf, text=txt, variable=v_type, value=val, bg=W, fg=BLK,
                           activebackground=W, selectcolor=GLT,
                           font=(F, 10, "bold")).pack(side=tk.LEFT, padx=(0, 20))

        lbl("Doc No *", 1)
        wrap_no = tk.Frame(fc, bg=G3)
        wrap_no.grid(row=1, column=1, padx=(2, 12), pady=(1, 6), sticky="ew")
        _no_vcmd = (self.register(lambda P: P.isdigit() or P == ""), "%P")
        e_no = tk.Entry(wrap_no, width=20, bg=G2, fg=BLK, insertbackground=BLK,
                        relief=tk.FLAT, font=(F, 10), bd=7,
                        validate="key", validatecommand=_no_vcmd)
        e_no.pack(fill=tk.X)
        _default_no = edit_row[2] if edit_row else next_no(default_type)
        _digits_only = re.sub(r"\D", "", _default_no)
        e_no.insert(0, _digits_only)

        lbl("Reference No", 1, 2)
        e_ref = ent(1, 3, 22, edit_row[15] if edit_row and len(edit_row) > 15 else "")

        def refresh_no():
            if not edit_row:
                new_val = re.sub(r"\D", "", next_no(v_type.get()))
                e_no.delete(0, tk.END); e_no.insert(0, new_val)
        for rb in tf.winfo_children(): rb.config(command=refresh_no)

        lbl("Subject *", 2)
        sw2 = tk.Frame(fc, bg=G3)
        sw2.grid(row=2, column=1, columnspan=3, padx=(2, 12), pady=(1, 6), sticky="ew")
        e_subj = tk.Entry(sw2, bg=G2, fg=BLK, insertbackground=BLK,
                          relief=tk.FLAT, font=(F, 10), bd=7)
        e_subj.pack(fill=tk.X)
        if edit_row: e_subj.insert(0, edit_row[3])

        lbl("Sender / Receiver *", 3); e_sr  = ent(3, 1, 26, edit_row[4] if edit_row else "")
        lbl("Department", 3, 2)
        v_dept = cbo(3, 3, ["", "Administration", "Finance", "HR", "IT", "Legal",
                             "Operations", "Procurement", "Sales", "Other"],
                    val=edit_row[5] if edit_row else "")

        lbl("Forwarded To", 4); e_fwd = ent(4, 1, 26, edit_row[6] if edit_row else "")
        lbl("Priority", 4, 2)
        v_pri = tk.StringVar(value=edit_row[10] if edit_row else "Normal")
        pf2 = tk.Frame(fc, bg=W)
        pf2.grid(row=4, column=3, sticky="w", padx=(2, 12), pady=(1, 6))
        for txt, val in [("Urgent", "Urgent"), ("Normal", "Normal"), ("Low", "Low")]:
            tk.Radiobutton(pf2, text=txt, variable=v_pri, value=val, bg=W, fg=BLK,
                           activebackground=W, selectcolor=GLT,
                           font=(F, 10)).pack(side=tk.LEFT, padx=(0, 14))

        lbl("Entry Date * (YYYY-MM-DD)", 5)

        def make_date_entry(row_n, col_n, default_val=""):
            wrap = tk.Frame(fc, bg=G3)
            wrap.grid(row=row_n, column=col_n, padx=(2, 12), pady=(1, 6), sticky="ew")
            vcmd = (self.register(
                lambda P: (len(P) <= 10 and all(ch.isdigit() or ch == "-" for ch in P))), "%P")
            e = tk.Entry(wrap, width=16, bg=G2, fg=BLK, insertbackground=BLK,
                         relief=tk.FLAT, font=(F, 10), bd=7,
                         validate="key", validatecommand=vcmd)
            e.pack(fill=tk.X)
            if default_val: e.insert(0, default_val)
            def on_focus_out(ev, entry=e):
                v = entry.get().strip()
                if v and len(v) == 8 and v.isdigit():
                    entry.delete(0, tk.END)
                    entry.insert(0, f"{v[:4]}-{v[4:6]}-{v[6:]}")
            e.bind("<FocusOut>", on_focus_out)
            return e

        e_edate = make_date_entry(5, 1, edit_row[7] if edit_row else tod())
        lbl("Doc Date", 5, 2)
        e_ddate = make_date_entry(5, 3, edit_row[8] if edit_row else tod())

        lbl("Due Date (YYYY-MM-DD)", 6)
        e_due = make_date_entry(6, 1, edit_row[9] if edit_row else "")
        lbl("Status", 6, 2)
        v_stat = cbo(6, 3, ["Pending", "In-Progress", "Completed", "Forwarded", "Archived"],
                    val=edit_row[11] if edit_row else "Pending")

        lbl("Tags", 7); e_tags = ent(7, 1, 26, edit_row[14] if edit_row and len(edit_row) > 14 else "")
        lbl("(e.g. tender, legal)", 7, 2)

        lbl("Remarks", 8)
        rw = tk.Frame(fc, bg=G3)
        rw.grid(row=8, column=1, columnspan=3, padx=(2, 12), pady=(1, 6), sticky="ew")
        e_rem = tk.Text(rw, height=3, bg=G2, fg=BLK, insertbackground=BLK,
                        font=(F, 10), relief=tk.FLAT, bd=7)
        e_rem.pack(fill=tk.X)
        if edit_row and edit_row[12]: e_rem.insert("1.0", edit_row[12])

        lbl("Attach File", 9)
        v_file = tk.StringVar(value=edit_row[13] if edit_row and len(edit_row) > 13 else "")
        firow = tk.Frame(fc, bg=W)
        firow.grid(row=9, column=1, columnspan=3, padx=(2, 12), pady=(1, 12), sticky="ew")
        fiwrap = tk.Frame(firow, bg=G3); fiwrap.pack(side=tk.LEFT, fill=tk.X, expand=True)
        fe_path = tk.Entry(fiwrap, textvariable=v_file, bg=G2, fg=BLK,
                           insertbackground=BLK, relief=tk.FLAT, font=(F, 9), bd=7, cursor="hand2")
        fe_path.pack(fill=tk.X)
        fe_path.bind("<Button-1>", lambda e: open_file(v_file.get()))
        mk_btn(firow, "Browse",
               lambda: v_file.set(filedialog.askopenfilename(
                   filetypes=[("All", "*.*"), ("PDF", "*.pdf")]) or v_file.get()),
               style="s", py=4, px=10).pack(side=tk.LEFT, padx=(6, 0))
        tk.Label(firow, text="(Click file name to open)", font=(F, 8), bg=W, fg=G4).pack(side=tk.LEFT, padx=8)

        old_status = edit_row[11] if edit_row else None

        def do_save():
            vals = {
                "doc_type":        v_type.get(),
                "doc_no":          e_no.get().strip(),
                "subject":         e_subj.get().strip(),
                "sender_receiver": e_sr.get().strip(),
                "department":      v_dept.get(),
                "forwarded_to":    e_fwd.get().strip(),
                "date_entry":      e_edate.get().strip(),
                "doc_date":        e_ddate.get().strip(),
                "due_date":        e_due.get().strip(),
                "priority":        v_pri.get(),
                "status":          v_stat.get(),
                "remarks":         e_rem.get("1.0", tk.END).strip(),
                "file_path":       v_file.get().strip(),
                "tags":            e_tags.get().strip(),
                "ref_no":          e_ref.get().strip(),
            }
            if not vals["doc_no"] or not vals["subject"] or not vals["sender_receiver"] or not vals["date_entry"]:
                messagebox.showwarning("Validation", "Please fill all required fields (*)"); return

            date_pattern = re.compile(r"^\d{4}-\d{2}-\d{2}$")
            for df, dlbl in [(vals["date_entry"], "Entry Date"), (vals["doc_date"], "Doc Date")]:
                if df and not date_pattern.match(df):
                    messagebox.showwarning("Invalid Date", f"{dlbl} must be in YYYY-MM-DD format\n(e.g. 2025-01-31)"); return
            if vals["due_date"] and not date_pattern.match(vals["due_date"]):
                messagebox.showwarning("Invalid Date", "Due Date must be in YYYY-MM-DD format\n(e.g. 2025-01-31)"); return

            try:
                c = con()
                if edit_row:
                    c.execute("""UPDATE documents SET
                        doc_type=?,doc_no=?,subject=?,sender_receiver=?,department=?,forwarded_to=?,
                        date_entry=?,doc_date=?,due_date=?,priority=?,status=?,remarks=?,
                        file_path=?,tags=?,ref_no=?,updated_at=? WHERE id=?""",
                        (vals["doc_type"], vals["doc_no"], vals["subject"], vals["sender_receiver"],
                         vals["department"], vals["forwarded_to"], vals["date_entry"], vals["doc_date"],
                         vals["due_date"], vals["priority"], vals["status"], vals["remarks"],
                         vals["file_path"], vals["tags"], vals["ref_no"], now(), edit_row[0]))
                    c.commit(); c.close()
                    audit(edit_row[0], "Updated", old_status, vals["status"],
                          "status", "Document edited", self.user["username"])
                else:
                    c.execute("""INSERT INTO documents
                        (doc_type,doc_no,subject,sender_receiver,department,forwarded_to,date_entry,
                         doc_date,due_date,priority,status,remarks,file_path,tags,ref_no,created_by,created_at)
                        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                        (vals["doc_type"], vals["doc_no"], vals["subject"], vals["sender_receiver"],
                         vals["department"], vals["forwarded_to"], vals["date_entry"], vals["doc_date"],
                         vals["due_date"], vals["priority"], vals["status"], vals["remarks"],
                         vals["file_path"], vals["tags"], vals["ref_no"], self.user["username"], now()))
                    nid = c.execute("SELECT last_insert_rowid()").fetchone()[0]
                    c.commit(); c.close()
                    audit(nid, "Created", None, vals["status"], "status",
                          "New document added", self.user["username"])
                self._go("docs")
            except sqlite3.IntegrityError:
                messagebox.showerror("Error", f"Doc No '{vals['doc_no']}' already exists!")

        bf = tk.Frame(body, bg=W); bf.pack(anchor="w", pady=(0, 10))
        mk_btn(bf, "  Save Document  ", do_save, style="p", side=tk.LEFT)
        mk_btn(bf, "  Cancel  ", lambda: self._go("docs"), style="s", side=tk.LEFT)

    # ════════════════════════════════════════════
    #  AUDIT HISTORY
    # ════════════════════════════════════════════
    def _pg_history(self):
        body, _ = self._page("Audit History", "Detailed action logs")
        hf = tk.Frame(body, bg=W, bd=1, relief=tk.SOLID); hf.pack(fill=tk.X, pady=(0, 10))
        ir = tk.Frame(hf, bg=W); ir.pack(fill=tk.X, padx=12, pady=8)
        tk.Label(ir, text="Filter by Doc ID / Doc No:", font=(F, 9, "bold"), bg=W, fg=BLK).pack(side=tk.LEFT)
        sw = tk.Frame(ir, bg=G3); sw.pack(side=tk.LEFT, padx=(4, 8))
        eh = tk.Entry(sw, width=22, bg=G2, fg=BLK, insertbackground=BLK,
                      relief=tk.FLAT, font=(F, 9), bd=6); eh.pack()
        frm, tv = mk_tree(body,
            ("ID", "Doc ID", "Action", "Field", "Old Value", "New Value", "Remarks", "By", "Date/Time"),
            [46, 56, 90, 80, 110, 110, 220, 110, 150], height=23)
        frm.pack(fill=tk.BOTH, expand=True)

        def load(*_):
            for r in tv.get_children(): tv.delete(r)
            kw = eh.get().strip()
            c = con(); cur = c.cursor()
            if kw:
                cur.execute("SELECT h.* FROM doc_history h JOIN documents d ON h.doc_id=d.id "
                            "WHERE CAST(h.doc_id AS TEXT) LIKE ? OR LOWER(d.doc_no) LIKE ? "
                            "ORDER BY h.id DESC",
                            (f"%{kw}%", f"%{kw.lower()}%"))
            else:
                cur.execute("SELECT * FROM doc_history ORDER BY id DESC LIMIT 600")
            for i, row in enumerate(cur.fetchall()):
                tv.insert("", "end", values=row, tags=("alt" if i % 2 else "",))
            c.close()

        mk_btn(ir, "Search",   load, style="p", py=4).pack(side=tk.LEFT)
        mk_btn(ir, "Show All", load, style="s", py=4).pack(side=tk.LEFT, padx=4)
        load()

    # ════════════════════════════════════════════
    #  REMINDERS
    # ════════════════════════════════════════════
    def _pg_reminders(self):
        body, _ = self._page("Reminders", "Due dates and urgent documents")
        c = con(); cur = c.cursor(); td = tod()
        cur.execute("SELECT doc_no,doc_type,subject,sender_receiver,due_date,status,priority "
                    "FROM documents WHERE due_date=? AND status NOT IN('Completed','Archived') "
                    "ORDER BY doc_no", (td,))
        due_today = cur.fetchall()
        cur.execute("SELECT doc_no,doc_type,subject,sender_receiver,due_date,status,priority "
                    "FROM documents WHERE priority='Urgent' AND status NOT IN('Completed','Archived') "
                    "ORDER BY date_entry DESC")
        urgent = cur.fetchall()
        cur.execute("SELECT doc_no,doc_type,subject,sender_receiver,due_date,status,priority "
                    "FROM documents WHERE due_date!='' AND due_date<? "
                    "AND status NOT IN('Completed','Archived') ORDER BY due_date", (td,))
        overdue = cur.fetchall()
        c.close()

        for sec_title, rows, bg, fnt in [
            (f"Due Today  ({len(due_today)} items)",   due_today, YLT,   (F, 9, "bold")),
            (f"Urgent Pending  ({len(urgent)} items)", urgent,    RLT,   (F, 9, "bold")),
            (f"Overdue  ({len(overdue)} items)",       overdue,   "#fde8e8" if not IS_DARK else RLT, (F, 9, "bold")),
        ]:
            lf = tk.LabelFrame(body, text=f"  {sec_title}  ", font=(F, 10, "bold"),
                               bg=W, fg=BLK, bd=1, relief=tk.SOLID)
            lf.pack(fill=tk.X, pady=(0, 12))
            if not rows:
                tk.Label(lf, text="  No items.", font=(F, 9), bg=W, fg=G4).pack(anchor="w", padx=12, pady=8)
            else:
                for row in rows:
                    rf = tk.Frame(lf, bg=bg); rf.pack(fill=tk.X, padx=8, pady=2)
                    tk.Label(rf, text=f"  {row[0]}  |  {row[2][:45]}  |  {row[3][:25]}"
                                      f"  |  Due: {row[4] or 'N/A'}  |  {row[5]}",
                             font=fnt, bg=bg, fg=BLK, anchor="w").pack(fill=tk.X, padx=6, pady=5)

    # ════════════════════════════════════════════
    #  REPORTS
    # ════════════════════════════════════════════
    def _pg_reports(self):
        body, _ = self._page("Reports & Export")
        exp_card = tk.Frame(body, bg=W, bd=1, relief=tk.SOLID)
        exp_card.pack(fill=tk.X, pady=(0, 16))
        tk.Label(exp_card, text="Export Options", font=(F, 11, "bold"),
                 bg=W, fg=BLK).pack(anchor="w", padx=16, pady=(14, 8))
        bf = tk.Frame(exp_card, bg=W); bf.pack(anchor="w", padx=16, pady=(0, 14))
        for txt, fn, st in [
            ("CSV Export",  lambda: self._csv(None), "p"),
            ("HTML Report", self._html,               "t"),
            ("TXT Report",  self._txt,                "s"),
            ("PDF Report",  self._pdf,                "d"),
        ]:
            mk_btn(bf, txt, fn, style=st, py=8, px=14, side=tk.LEFT)

        tk.Label(body, text="Summary by Type / Status / Priority",
                 font=(F, 11, "bold"), bg=W, fg=BLK).pack(anchor="w", pady=(0, 6))
        frm, tv = mk_tree(body, ("Doc Type", "Status", "Priority", "Count"),
                          [120, 130, 100, 80], height=15)
        frm.pack(fill=tk.BOTH, expand=True)
        c = con(); cur = c.cursor()
        cur.execute("SELECT doc_type,status,priority,COUNT(*) FROM documents "
                    "GROUP BY doc_type,status,priority ORDER BY doc_type,status")
        for i, row in enumerate(cur.fetchall()):
            tv.insert("", "end", values=row,
                      tags=(("IN" if row[0] == "INWARD" else "OUT"), "alt" if i % 2 else ""))
        c.close()

    def _all_docs(self, ft=None):
        c = con(); cur = c.cursor()
        if ft: cur.execute("SELECT * FROM documents WHERE doc_type=? ORDER BY id DESC", (ft,))
        else:  cur.execute("SELECT * FROM documents ORDER BY id DESC")
        rows = cur.fetchall(); hdrs = [d[0] for d in cur.description]; c.close()
        return hdrs, rows

    def _csv(self, ft=None):
        hdrs, rows = self._all_docs(ft)
        if not rows: messagebox.showinfo("Empty", "No data to export."); return
        path = filedialog.asksaveasfilename(defaultextension=".csv",
            initialfile=f"DMS_{ft or 'ALL'}_{tod()}.csv", filetypes=[("CSV", "*.csv")])
        if not path: return
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f); w.writerow(hdrs); w.writerows(rows)
        if messagebox.askyesno("Done", f"{len(rows)} records exported!\nOpen file?"): open_file(path)

    def _txt(self):
        hdrs, rows = self._all_docs()
        if not rows: messagebox.showinfo("Empty", "No data to export."); return
        path = filedialog.asksaveasfilename(defaultextension=".txt",
            initialfile=f"DMS_Report_{tod()}.txt", filetypes=[("Text", "*.txt")])
        if not path: return
        with open(path, "w", encoding="utf-8") as f:
            f.write(f"DMS EXPORT — {now()}\nRecords: {len(rows)}\n\n")
            for row in rows:
                f.write("-" * 50 + "\n")
                for h, v in zip(hdrs, row): f.write(f"{h}: {v or ''}\n")
        if messagebox.askyesno("Done", f"{len(rows)} records exported!\nOpen?"): open_file(path)

    def _html(self):
        hdrs, rows = self._all_docs()
        if not rows: messagebox.showinfo("Empty", "No data."); return
        path = filedialog.asksaveasfilename(defaultextension=".html",
            initialfile=f"DMS_Report_{tod()}.html", filetypes=[("HTML", "*.html")])
        if not path: return
        th = "".join(f"<th>{h}</th>" for h in hdrs)
        tbody = "".join(f"<tr>{''.join(f'<td>{v or str()}</td>' for v in r)}</tr>\n" for r in rows)
        html = (f"<html><body><h2>DMS Report {now()}</h2>"
                f"<table border='1'><thead><tr>{th}</tr></thead><tbody>{tbody}</tbody></table></body></html>")
        with open(path, "w", encoding="utf-8") as f: f.write(html)
        if messagebox.askyesno("Done", "Exported!\nOpen in browser?"):
            webbrowser.open(f"file:///{os.path.abspath(path)}")

    def _pdf(self):
        try:
            from fpdf import FPDF
        except ImportError:
            messagebox.showwarning("Library Required",
                "To export to PDF, please install the library by running this in your command prompt:\n\npip install fpdf")
            return
        hdrs, rows = self._all_docs()
        if not rows: messagebox.showinfo("Empty", "No data."); return
        path = filedialog.asksaveasfilename(defaultextension=".pdf",
            initialfile=f"DMS_Report_{tod()}.pdf", filetypes=[("PDF", "*.pdf")])
        if not path: return
        pdf = FPDF(); pdf.add_page(); pdf.set_font("Arial", size=10)
        pdf.cell(200, 10, txt=f"DMS Document Report - {now()}", ln=True, align="C")
        pdf.ln(10)
        for row in rows:
            pdf.set_font("Arial", "B", 10)
            pdf.cell(0, 8, txt=f"Doc No: {row[2]}  |  Type: {row[1]}", ln=True)
            pdf.set_font("Arial", "", 9)
            pdf.cell(0, 6, txt=f"Subject: {row[3]}", ln=True)
            pdf.cell(0, 6, txt=f"Sender/Receiver: {row[4]} | Status: {row[11]}", ln=True)
            pdf.cell(0, 6, txt="-" * 80, ln=True)
        pdf.output(path)
        if messagebox.askyesno("Done", "PDF saved!\nOpen file?"): open_file(path)

    # ════════════════════════════════════════════
    #  USER MANAGEMENT  (Admin only)
    # ════════════════════════════════════════════
    def _pg_user_mgmt(self):
        body, hdr = self._page("User Management", "Manage system users (Admin only)")
        mk_btn(hdr, "+ Add New User", lambda: self._um_add(), style="p", py=5).pack(side=tk.RIGHT, padx=(4, 0))

        fb = tk.Frame(body, bg=G1, bd=1, relief=tk.SOLID); fb.pack(fill=tk.X, pady=(0, 10))
        fbr = tk.Frame(fb, bg=G1); fbr.pack(fill=tk.X, padx=10, pady=7)

        tk.Label(fbr, text="Search:", font=(F, 9, "bold"), bg=G1, fg=BLK).pack(side=tk.LEFT)
        self._us = tk.StringVar()
        self._us.trace("w", lambda *a: self._um_load())
        sw = tk.Frame(fbr, bg=G3); sw.pack(side=tk.LEFT, padx=(4, 10))
        tk.Entry(sw, textvariable=self._us, width=20, bg=G2, fg=BLK,
                 insertbackground=BLK, relief=tk.FLAT, font=(F, 9), bd=6).pack()

        tk.Label(fbr, text="Role:", font=(F, 9, "bold"), bg=G1, fg=BLK).pack(side=tk.LEFT)
        self._ur = tk.StringVar(value="ALL")
        rcb = ttk.Combobox(fbr, textvariable=self._ur,
                           values=["ALL", "Admin", "Clerk", "User", "Manager"],
                           width=10, font=(F, 9), state="readonly")
        rcb.pack(side=tk.LEFT, padx=(4, 8), ipady=3)
        rcb.bind("<<ComboboxSelected>>", lambda e: self._um_load())

        tk.Label(fbr, text="Status:", font=(F, 9, "bold"), bg=G1, fg=BLK).pack(side=tk.LEFT)
        self._ust = tk.StringVar(value="ALL")
        scb = ttk.Combobox(fbr, textvariable=self._ust,
                           values=["ALL", "Active", "Inactive"],
                           width=10, font=(F, 9), state="readonly")
        scb.pack(side=tk.LEFT, padx=(4, 8), ipady=3)
        scb.bind("<<ComboboxSelected>>", lambda e: self._um_load())

        ab = tk.Frame(body, bg=G1); ab.pack(side=tk.BOTTOM, fill=tk.X)
        self._usl = tk.Label(ab, text="", font=(F, 9), bg=ACC, fg="#ffffff", anchor="w")
        self._usl.pack(side=tk.LEFT, fill=tk.X, expand=True, ipady=5, padx=6)
        mk_btn(ab, "Reset Password", self._um_reset_pw,   style="p", py=4, px=10, side=tk.RIGHT)
        mk_btn(ab, "Deactivate",     self._um_deactivate, style="p", py=4, px=10, side=tk.RIGHT)
        mk_btn(ab, "Activate",       self._um_activate,   style="p", py=4, px=10, side=tk.RIGHT)
        mk_btn(ab, "Edit User",      self._um_edit,       style="p", py=4, px=10, side=tk.RIGHT)
        mk_btn(ab, "🗑 Delete User",  self._um_delete,     style="d", py=4, px=10, side=tk.RIGHT)

        cols   = ("ID", "Username", "Full Name", "Role", "Department", "Email", "Status", "Last Login", "Created")
        widths = [36, 100, 140, 70, 110, 160, 60, 130, 110]
        tf, self._utv = mk_tree(body, cols, widths)
        tf.pack(fill=tk.BOTH, expand=True, pady=(0, 4))
        self._utv.bind("<<TreeviewSelect>>", self._um_sel)
        self._usel = None
        self._um_load()

    def _um_load(self):
        if not hasattr(self, "_utv"): return
        for r in self._utv.get_children(): self._utv.delete(r)
        srch = getattr(self, "_us",  tk.StringVar()).get().strip().lower()
        role = getattr(self, "_ur",  tk.StringVar(value="ALL")).get()
        stat = getattr(self, "_ust", tk.StringVar(value="ALL")).get()

        c = con(); cur = c.cursor()
        q = "SELECT id,username,full_name,role,department,email,is_active,last_login,created_at FROM users"
        conds = []; params = []
        if role != "ALL": conds.append("role=?"); params.append(role)
        if stat == "Active":   conds.append("is_active=1")
        if stat == "Inactive": conds.append("is_active=0")
        if srch:
            conds.append("(LOWER(username) LIKE ? OR LOWER(full_name) LIKE ? OR LOWER(email) LIKE ?)")
            params += [f"%{srch}%"] * 3
        if conds: q += " WHERE " + " AND ".join(conds)
        q += " ORDER BY id"
        cur.execute(q, params); rows = cur.fetchall(); c.close()

        for i, row in enumerate(rows):
            status_txt = "✅ Active" if row[6] else "❌ Inactive"
            display = (i + 1, row[1], row[2] or "—", row[3],
                       row[4] or "—", row[5] or "—",
                       status_txt, row[7] or "Never", row[8])
            tag = "alt" if i % 2 else ""
            if not row[6]: tag = "ovr"
            self._utv.insert("", "end", iid=str(row[0]), values=display, tags=(tag,))

        total    = len(rows)
        active   = sum(1 for r in rows if r[6])
        inactive = total - active
        if hasattr(self, "_usl"):
            self._usl.config(text=f"  Total: {total}  |  Active: {active}  |  Inactive: {inactive}")

    def _um_sel(self, _):
        sel = self._utv.selection()
        if sel: self._usel = int(sel[0])

    def _um_need_sel(self):
        if not self._usel:
            messagebox.showwarning("Select", "Please select a user first!"); return False
        return True

    def _um_add(self):
        UserDlg(self, None, self._um_load)

    def _um_edit(self):
        if not self._um_need_sel(): return
        c = con(); cur = c.cursor()
        cur.execute("SELECT * FROM users WHERE id=?", (self._usel,))
        row = cur.fetchone(); c.close()
        if row: UserEditDlg(self, row, self._um_load, current_user=self.user)

    def _um_deactivate(self):
        if not self._um_need_sel(): return
        c = con(); cur = c.cursor()
        cur.execute("SELECT username,role FROM users WHERE id=?", (self._usel,))
        row = cur.fetchone(); c.close()
        if not row: return
        if row[1] == "Admin" and row[0] == self.user["username"]:
            messagebox.showwarning("Cannot Deactivate", "You cannot deactivate your own account!"); return
        if not messagebox.askyesno("Confirm", f"Deactivate user '{row[0]}'?\nThey will not be able to login."): return
        c = con(); c.execute("UPDATE users SET is_active=0 WHERE id=?", (self._usel,))
        c.commit(); c.close()
        messagebox.showinfo("Done", f"User '{row[0]}' deactivated.")
        self._usel = None; self._um_load()

    def _um_activate(self):
        if not self._um_need_sel(): return
        c = con(); cur = c.cursor()
        cur.execute("SELECT username FROM users WHERE id=?", (self._usel,))
        row = cur.fetchone(); c.close()
        if not row: return
        if not messagebox.askyesno("Confirm", f"Activate user '{row[0]}'?"): return
        c = con(); c.execute("UPDATE users SET is_active=1 WHERE id=?", (self._usel,))
        c.commit(); c.close()
        messagebox.showinfo("Done", f"User '{row[0]}' activated.")
        self._usel = None; self._um_load()

    def _um_reset_pw(self):
        if not self._um_need_sel(): return
        c = con(); cur = c.cursor()
        cur.execute("SELECT username FROM users WHERE id=?", (self._usel,))
        row = cur.fetchone(); c.close()
        if not row: return
        new_pw = simpledialog.askstring("Reset Password",
            f"Enter new password for '{row[0]}':\n(Minimum 6 characters)",
            show="*", parent=self)
        if not new_pw: return
        if len(new_pw) < 6:
            messagebox.showwarning("Too Short", "Password must be at least 6 characters!"); return
        c = con(); c.execute("UPDATE users SET password=? WHERE id=?", (hp(new_pw), self._usel))
        c.commit(); c.close()
        messagebox.showinfo("Done", f"Password for '{row[0]}' has been reset successfully.")

    def _um_delete(self):
        if not self._um_need_sel(): return
        c = con(); cur = c.cursor()
        cur.execute("SELECT username, role FROM users WHERE id=?", (self._usel,))
        row = cur.fetchone(); c.close()
        if not row: return
        if row[0] == self.user["username"]:
            messagebox.showwarning("Cannot Delete", "You cannot delete your own account!"); return
        if not messagebox.askyesno("Confirm Delete",
            f"Permanently delete user '{row[0]}'?\n\nThis action CANNOT be undone!",
            icon="warning"): return
        c = con(); c.execute("DELETE FROM users WHERE id=?", (self._usel,))
        c.commit(); c.close()
        messagebox.showinfo("Deleted", f"User '{row[0]}' has been permanently deleted.")
        self._usel = None; self._um_load()

    def _um_export_xlsx(self):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            ans = messagebox.askyesno("Library Required",
                "To export to Excel, the 'openpyxl' library is required.\n\n"
                "Install it by running:\n\n    pip install openpyxl\n\n"
                "Would you like to export as CSV instead?")
            if ans: self._um_export_csv()
            return

        c = con(); cur = c.cursor()
        cur.execute("SELECT id, username, full_name, email, phone, department, "
                    "role, is_active, last_login, created_at FROM users ORDER BY id")
        rows = cur.fetchall(); c.close()
        if not rows: messagebox.showinfo("Empty", "No users to export."); return

        path = filedialog.asksaveasfilename(defaultextension=".xlsx",
            initialfile=f"Users_Export_{tod()}.xlsx",
            filetypes=[("Excel Workbook", "*.xlsx")])
        if not path: return

        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Users"
        header_font  = Font(name="Segoe UI", bold=True, color="FFFFFF", size=11)
        header_fill  = PatternFill("solid", fgColor="1A6B4A")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border  = Border(left=Side(style="thin"), right=Side(style="thin"),
                              top=Side(style="thin"),  bottom=Side(style="thin"))
        alt_fill = PatternFill("solid", fgColor="F0FAF4")

        headers    = ["ID", "Username", "Full Name", "Email", "Phone",
                      "Department", "Role", "Status", "Last Login", "Created At"]
        col_widths = [6, 18, 22, 28, 16, 18, 12, 10, 22, 22]

        for ci, (h, w) in enumerate(zip(headers, col_widths), start=1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.font = header_font; cell.fill = header_fill
            cell.alignment = header_align; cell.border = thin_border
            ws.column_dimensions[cell.column_letter].width = w
        ws.row_dimensions[1].height = 28

        data_font  = Font(name="Segoe UI", size=10)
        data_align = Alignment(vertical="center", wrap_text=False)
        for ri, row in enumerate(rows, start=2):
            status = "Active" if row[7] else "Inactive"
            values = [row[0], row[1], row[2] or "", row[3] or "",
                      row[4] or "", row[5] or "", row[6],
                      status, row[8] or "Never", row[9] or ""]
            fill = alt_fill if ri % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
            for ci, val in enumerate(values, start=1):
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.font = data_font; cell.alignment = data_align
                cell.border = thin_border; cell.fill = fill
            ws.row_dimensions[ri].height = 20

        ws.freeze_panes = "A2"
        wb.save(path)
        if messagebox.askyesno("Export Complete",
                               f"{len(rows)} user(s) exported successfully!\n\nOpen file?"):
            open_file(path)

    def _um_export_csv(self):
        import csv as _csv
        c = con(); cur = c.cursor()
        cur.execute("SELECT id, username, full_name, email, phone, department, "
                    "role, is_active, last_login, created_at FROM users ORDER BY id")
        rows = cur.fetchall(); c.close()
        if not rows: messagebox.showinfo("Empty", "No users to export."); return
        path = filedialog.asksaveasfilename(defaultextension=".csv",
            initialfile=f"Users_Export_{tod()}.csv", filetypes=[("CSV", "*.csv")])
        if not path: return
        headers = ["ID", "Username", "Full Name", "Email", "Phone",
                   "Department", "Role", "Status", "Last Login", "Created At"]
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = _csv.writer(f); w.writerow(headers)
            for row in rows:
                status = "Active" if row[7] else "Inactive"
                w.writerow([row[0], row[1], row[2] or "", row[3] or "",
                             row[4] or "", row[5] or "", row[6],
                             status, row[8] or "Never", row[9] or ""])
        if messagebox.askyesno("Done", f"{len(rows)} users exported!\nOpen file?"): open_file(path)

    # ════════════════════════════════════════════
    #  SETTINGS
    # ════════════════════════════════════════════
    def _pg_settings(self):
        body, _ = self._page("Settings", "Application preferences")
        c = con(); cur = c.cursor()
        cur.execute("SELECT key,value FROM settings"); saved = dict(cur.fetchall()); c.close()

        fc = tk.Frame(body, bg=W, bd=1, relief=tk.SOLID); fc.pack(fill=tk.X)
        fc.columnconfigure(0, weight=1)

        def sf(lbl, key, default, row):
            tk.Label(fc, text=lbl, font=(F, 10, "bold"), bg=W, fg=BLK).grid(
                row=row * 2, column=0, sticky="w", padx=16, pady=(12, 1))
            wrap = tk.Frame(fc, bg=G3); wrap.grid(row=row * 2 + 1, column=0, sticky="ew", padx=16, pady=(0, 6))
            e = tk.Entry(wrap, bg=G2, fg=BLK, insertbackground=BLK,
                         relief=tk.FLAT, font=(F, 10), bd=7); e.pack(fill=tk.X)
            e.insert(0, saved.get(key, default)); return e, key

        def cbof(lbl, key, vals, default, row):
            tk.Label(fc, text=lbl, font=(F, 10, "bold"), bg=W, fg=BLK).grid(
                row=row * 2, column=0, sticky="w", padx=16, pady=(12, 1))
            wrap = tk.Frame(fc, bg=G3); wrap.grid(row=row * 2 + 1, column=0, sticky="ew", padx=16, pady=(0, 6))
            v = tk.StringVar(value=saved.get(key, default))
            cb = ttk.Combobox(wrap, textvariable=v, values=vals, font=(F, 10), state="readonly")
            cb.pack(fill=tk.X, ipady=5); return v, key

        e1, k1 = sf("Organization Name", "org_name", "My Organization", 0)
        v2, k2 = cbof("Enable Dark Mode (Requires restart to apply)", "dark_mode", ["Yes", "No"], "No", 1)
        v3, k3 = cbof("Default Export Format", "default_export", ["CSV", "PDF", "HTML", "TXT"], "CSV", 2)
        v4, k4 = cbof("Dashboard Table Rows limit", "dash_rows", ["10", "20", "50", "100"], "10", 3)
        fields = [(e1, k1), (v2, k2), (v3, k3), (v4, k4)]

        msg = tk.StringVar()
        tk.Label(fc, textvariable=msg, font=(F, 9), bg=W, fg=ACC).grid(row=9, column=0, pady=(4, 0), padx=16)

        def save():
            c = con()
            for e, key in fields:
                val = e.get().strip() if hasattr(e, "get") else e
                c.execute("INSERT OR REPLACE INTO settings(key,value) VALUES(?,?)", (key, val))
            c.commit(); c.close()
            msg.set("Settings saved successfully! ✓ (Restart app to apply theme changes)")

        mk_btn(fc, "  Save Settings  ", save, style="p", py=9, px=0).grid(
            row=10, column=0, sticky="ew", padx=16, pady=(10, 16))

        divider(body, 18)
        tk.Label(body, text="💡 Suggest a New Feature",
                 font=(F, 12, "bold"), bg=W, fg=BLK).pack(anchor="w", pady=(0, 6))
        sug_fr = tk.Frame(body, bg=W, bd=1, relief=tk.SOLID); sug_fr.pack(fill=tk.X)
        tk.Label(sug_fr,
                 text="Have an idea to improve this system? Write it below and save it for the developer.",
                 font=(F, 9), bg=W, fg=G5).pack(anchor="w", padx=16, pady=(10, 4))
        sug_wrap = tk.Frame(sug_fr, bg=G3); sug_wrap.pack(fill=tk.X, padx=16, pady=(2, 10))
        sug_txt = tk.Text(sug_wrap, height=4, bg=G2, fg=BLK, insertbackground=BLK,
                          font=(F, 10), relief=tk.FLAT, bd=7)
        sug_txt.pack(fill=tk.X)

        c2 = con(); cur2 = c2.cursor()
        cur2.execute("SELECT value FROM settings WHERE key='feature_suggestions'")
        existing_sug = cur2.fetchone(); c2.close()
        if existing_sug and existing_sug[0]: sug_txt.insert("1.0", existing_sug[0])

        sug_msg = tk.StringVar()
        tk.Label(sug_fr, textvariable=sug_msg, font=(F, 9), bg=W, fg=ACC).pack(padx=16)

        def save_suggestion():
            text = sug_txt.get("1.0", tk.END).strip()
            if not text: sug_msg.set("Please write your suggestion first."); return
            c2 = con()
            c2.execute("INSERT OR REPLACE INTO settings(key,value) VALUES(?,?)",
                       ("feature_suggestions", text))
            c2.commit(); c2.close()
            sug_msg.set("✓ Suggestion saved successfully!")

        mk_btn(sug_fr, "  Save Suggestion  ", save_suggestion, style="t", py=7, px=16).pack(
            anchor="w", padx=16, pady=(4, 14))

    # ════════════════════════════════════════════
    #  BACKUP / RESTORE
    # ════════════════════════════════════════════
    def _pg_backup(self):
        from config import DB
        body, _ = self._page("Backup & Restore")

        def section(title, desc, btn_text, cmd, sty):
            f = tk.Frame(body, bg=W, bd=1, relief=tk.SOLID); f.pack(fill=tk.X, pady=(0, 12))
            tk.Label(f, text=title, font=(F, 12, "bold"), bg=W, fg=BLK).pack(anchor="w", padx=16, pady=(12, 2))
            tk.Label(f, text=desc, font=(F, 9), bg=W, fg=G5).pack(anchor="w", padx=16)
            mk_btn(f, f"  {btn_text}  ", cmd, style=sty, py=8, px=16).pack(anchor="w", padx=16, pady=10)

        def do_bk():
            ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            path = filedialog.asksaveasfilename(defaultextension=".db",
                initialfile=f"DMS_Backup_{ts}.db", filetypes=[("Database", "*.db")])
            if not path: return
            shutil.copy2(DB, path)
            messagebox.showinfo("Done", f"Backup saved!\n{path}")

        def do_rs():
            path = filedialog.askopenfilename(filetypes=[("Database", "*.db")])
            if not path: return
            if not messagebox.askyesno("Confirm", "Current data will be replaced!\nConfirm?"): return
            shutil.copy2(path, DB)
            messagebox.showinfo("Done", "Restore completed! Please restart the app.")

        section("Create Backup", "Save database to a .db file.", "Create Backup", do_bk, "p")
        section("Restore from Backup", "WARNING: Current data will be replaced.", "Restore Backup", do_rs, "d")

    # ────────────────────────────────────────────
    #  MISC
    # ────────────────────────────────────────────
    def _upd_alerts(self):
        pass  # Alerts feature removed

    def _alerts_popup(self):
        pass  # Alerts feature removed

    def _logout(self):
        if messagebox.askyesno("Logout", "Are you sure you want to log out?"):
            self.destroy()
            self.login.deiconify()
            self.login.state("zoomed")
            self.login._pw.delete(0, tk.END)
            self.login._msg.set("")
            self.login._u.focus_set()

    def _quit(self):
        if messagebox.askyesno("Exit", "Do you want to close the application?"):
            self.login.destroy()
